from __future__ import annotations

import logging
from dataclasses import dataclass

from app.services.pdf_parser_service import ParsedPdf, PdfPage

logger = logging.getLogger(__name__)


class DocxParseError(Exception):
    """Raised when a Word document cannot be read or parsed."""


@dataclass
class DocxParserService:
    """Parse Word documents (.docx) into ParsedPdf format."""

    def parse(self, file_path: str) -> ParsedPdf:
        """Parse a .docx file and return ParsedPdf."""
        try:
            import docx
        except ImportError:
            raise DocxParseError(
                "python-docx is required for Word document support. "
                "Install it with: pip install python-docx"
            )

        try:
            doc = docx.Document(file_path)
        except Exception as exc:
            raise DocxParseError(f"Failed to read Word document: {exc}") from exc

        pages: list[PdfPage] = []
        current_paragraphs: list[str] = []
        page_num = 1

        # Extract text from paragraphs
        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            # If paragraph is a major heading (Heading 1 or Heading 2), treat as section boundary
            style_name = ""
            try:
                if p.style and hasattr(p.style, "name") and p.style.name:
                    style_name = p.style.name.lower()
            except Exception:
                pass

            if ("heading 1" in style_name or "heading 2" in style_name) and current_paragraphs:
                pages.append(PdfPage(
                    page_number=page_num,
                    extracted_text="\n\n".join(current_paragraphs),
                ))
                page_num += 1
                current_paragraphs = []
            current_paragraphs.append(text)

        # Extract text from tables
        for table in doc.tables:
            table_lines = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
                if cells:
                    table_lines.append(" | ".join(cells))
            if table_lines:
                current_paragraphs.append("\n".join(table_lines))

        # Save remaining content
        if current_paragraphs:
            pages.append(PdfPage(
                page_number=page_num,
                extracted_text="\n\n".join(current_paragraphs),
            ))

        if not pages:
            raise DocxParseError("No readable text or tables found in Word document")

        return ParsedPdf(pages=pages, total_pages=len(pages))
